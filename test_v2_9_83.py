"""2.9.83: Art. 31 deduction ratio of each quarter on its own (owner decision), and the audit report pack:
company / auditor information entered once for every year, centred cover page, Excel layout and extra notes."""
import tests_setup  # noqa: F401  the sample company of the window tests
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

import vat_return
from database import Database


class _Book(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.addCleanup(self.folder.cleanup)
        self.db = Database(Path(self.folder.name) / "books.db"); self.db.initialize("secret12345"); self.addCleanup(self.db.release)
        self.user = 1
        self.db.save_exchange_rate({"date_from": "01-01-2025", "date_to": "31-12-2025", "from_currency": "USD", "to_currency": "LBP", "rate": "89500"}, self.user)

    def sale(self, date, amount, treatment="standard", rate=11):
        return self.db.create_manual_invoice({"invoice_date": date, "party_name": "Client", "kind": "sales", "currency": "USD", "status": "posted", "vat_treatment": treatment},
                                             [{"description": "S", "quantity": 1, "unit_price": amount, "vat_rate": rate}], self.user)

    def purchase(self, date, amount, use="mixed"):
        return self.db.create_manual_invoice({"invoice_date": date, "party_name": "Supplier", "kind": "purchases", "currency": "USD", "status": "posted", "vat_use": use},
                                             [{"description": "P", "quantity": 1, "unit_price": amount, "vat_rate": 11}], self.user)


class QuarterRatioTest(_Book):
    def test_each_quarter_uses_its_own_turnover(self):
        self.assertEqual(vat_return.vat_ratio_method(self.db), "quarter")  # the default
        self.sale("10-02-2025", 8000); self.sale("11-02-2025", 2000, "exempt", 0)  # Q1: 80%
        self.sale("10-05-2025", 5000); self.sale("11-05-2025", 5000, "exempt", 0)  # Q2: 50%
        self.purchase("15-05-2025", 1000)
        q1 = vat_return.build_vat_return(self.db, 2025, 1); q2 = vat_return.build_vat_return(self.db, 2025, 2)
        self.assertEqual(q1["deduction_ratio"], Decimal("0.8")); self.assertEqual(q2["deduction_ratio"], Decimal("0.5"))
        self.assertIn("Q2 2025 only", q2["ratio_source"]); self.assertEqual(q2["ratio_method"], "quarter")
        self.assertEqual(q2["per_currency"]["USD"]["prorata"]["vat"], Decimal("-55.00"))  # 50% of the 110 mixed-use VAT
        self.assertEqual(q2["ytd_turnover_lbp"]["taxable"], 5000 * 89500)  # the quarter's turnover, not the year's

    def test_provisional_ratio_is_not_used_and_q4_has_no_annual_adjustment(self):
        self.db.save_vat_provisional_ratio(2025, "100", self.user)
        self.sale("10-02-2025", 5000); self.purchase("15-02-2025", 1000)
        vat_return.save_return(self.db, 2025, 1, self.user)
        self.sale("10-11-2025", 5000, "exempt", 0); self.sale("11-11-2025", 5000)
        q4 = vat_return.build_vat_return(self.db, 2025, 4)
        self.assertEqual(q4["deduction_ratio"], Decimal("0.5"))  # Q4 sales only
        self.assertEqual(q4["totals_lbp"]["annual_adjustment"], 0); self.assertEqual(q4["annual_adjustment_detail"], [])

    def test_the_annual_method_can_still_be_chosen(self):
        self.db.save_settings({"vat_ratio_method": "annual"}, self.user)
        self.assertEqual(vat_return.vat_ratio_method(self.db), "annual")
        self.db.save_vat_provisional_ratio(2025, "90", self.user); self.sale("10-02-2025", 1000)
        self.assertEqual(vat_return.build_vat_return(self.db, 2025, 1)["deduction_ratio"], Decimal("0.9"))
        with self.assertRaisesRegex(ValueError, "quarter or annual"): self.db.save_settings({"vat_ratio_method": "monthly"}, self.user)


class AuditPackTest(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.addCleanup(self.folder.cleanup)
        self.dbs = {}
        for year in (2024, 2025):
            db = Database(Path(self.folder.name) / f"{year}.db"); db.initialize("secret12345"); self.addCleanup(db.release); self.dbs[year] = db
            db.save_journal_voucher(dict(entry_date=f"01-03-{year}", currency="USD", description="sale"),
                                    [dict(account_code="531", debit="1000", credit=0), dict(account_code="7011", debit=0, credit="1000")], 1)

    def test_information_entered_once_reaches_every_year_and_every_text(self):
        import financial_statements as fs
        cfg = {"basis": "USD", "info": {"legal_form": "S.A.L.", "cr_number": "98765", "auditor_firm": "Saber for Audit", "report_city": "Tripoli"},
               "supplements": {"cf_operating": "5"}}
        fs.save_config(self.dbs[2025], cfg, 1, list(self.dbs.values()))
        other = fs.config(self.dbs[2024])
        self.assertEqual(other["info"]["cr_number"], "98765")
        self.assertNotIn("supplements", other)  # the cash flow amounts stay with their own year
        pack = fs.build({2024: self.dbs[2024]}, {"years": "2024", "basis": "USD"})
        text = str(pack["sections"])
        self.assertIn("is a S.A.L. registered in Lebanon under commercial register number 98765", text)
        self.assertIn("Saber for Audit", text); self.assertIn("Tripoli, [Date of the auditor's report]", text)
        self.assertIn("Auditor: Saber for Audit", pack["meta"])
        self.assertIn("Company & auditor information not entered", text)  # incorporation date etc. still to complete

    def test_an_older_year_without_settings_takes_them_from_another_year(self):
        import financial_statements as fs
        with self.dbs[2025].connect() as c:
            c.execute("INSERT INTO app_settings(key,value) VALUES(?,?)", (fs.KEY, '{"info": {"auditor_firm": "Firm A"}, "notes": {"Entity and activities": "Own text"}}'))
        self.assertEqual(fs.config(self.dbs[2024], list(self.dbs.values()))["info"]["auditor_firm"], "Firm A")
        self.assertIn("Own text", str(fs.build(self.dbs, {"years": "2024,2025"})["sections"]))

    def test_standard_text_saved_by_an_older_version_follows_the_new_one(self):
        import financial_statements as fs
        old = next(t for t in fs.PREVIOUS_DEFAULTS if t.startswith("{company}"))
        fs.save_config(self.dbs[2025], {"notes": {"Entity and activities": old}, "info": {"legal_form": "S.A.R.L."}}, 1)
        self.assertIn("is a S.A.R.L. registered", str(fs.build({2025: self.dbs[2025]}, {"years": "2025"})["sections"]))

    def test_cover_page_and_excel_sheets(self):
        import financial_statements as fs
        from openpyxl import load_workbook
        from pypdf import PdfReader
        from report_export import export_sections_excel, export_sections_pdf
        fs.save_config(self.dbs[2025], {"info": {"auditor_firm": "Saber for Audit"}}, 1, list(self.dbs.values()))
        pack = fs.build(self.dbs, {"years": "2024,2025", "basis": "USD"})
        pdf = Path(self.folder.name) / "pack.pdf"; xlsx = Path(self.folder.name) / "pack.xlsx"
        export_sections_pdf(pdf, pack["title"], pack["meta"], pack["sections"]); export_sections_excel(xlsx, pack["title"], pack["meta"], pack["sections"])
        cover = PdfReader(pdf).pages[0].extract_text() + PdfReader(pdf).pages[1].extract_text()  # 2.9.100: contents on page 2
        for words in ("FINANCIAL STATEMENTS", "CONTENTS", "Statement of Cash Flows", "Saber for Audit"): self.assertIn(words, cover)
        self.assertNotIn("OPINION", cover)  # the auditor's report starts on the next page
        wb = load_workbook(xlsx)
        self.assertEqual(wb.sheetnames, ["Cover", "Auditor's Report", "Financial Position", "Profit or Loss", "Changes in Equity", "Cash Flows", "Notes", "Review Points"])
        position = wb["Financial Position"]
        self.assertTrue(position.page_setup.fitToWidth); self.assertEqual(position.page_setup.orientation, "portrait")
        amounts = [c for row in position.iter_rows() for c in row if isinstance(c.value, (int, float))]
        self.assertTrue(amounts and all("(" in c.number_format for c in amounts))  # negatives in brackets
        wb.close()


if __name__ == "__main__":
    unittest.main()
