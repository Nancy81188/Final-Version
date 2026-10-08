import tests_setup  # noqa: F401  2.9.74: the sample company of the tests (a new installation has none)
import tempfile
import unittest
from pathlib import Path
from decimal import Decimal
from database import Database
import financial_statements as fs

class FinancialStatementsTest(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.addCleanup(self.folder.cleanup)
        self.db=Database(Path(self.folder.name)/'2025.db'); self.db.initialize('secret12345')
    def voucher(self, amount, debit='531', credit='7011', date='01-03-2025', **extra):
        return self.db.save_journal_voucher(dict(entry_date=date,currency='USD',description='FS test',**extra),
            [dict(account_code=debit,debit=str(amount),credit=0),dict(account_code=credit,debit=0,credit=str(amount))],1)
    def test_profit_opening_closing_and_balance_equation(self):
        self.voucher(500,credit='101',date='01-01-2025',voucher_type='04')
        self.voucher(1000); self.voucher(200,debit='6311',credit='531')
        # Simulated year-close transfer must not zero profit or double equity.
        closing=self.voucher(1000,debit='7011',credit='138',date='31-12-2025',voucher_type='05',entry_number='CLOSE-FS')
        with self.db.connect() as conn: conn.execute("UPDATE journal_entries SET description='CLOSING 6&7 - 2025' WHERE id=?",(closing["voucher"]["id"],))
        d=fs.year_data(self.db,2025,'USD')
        self.assertEqual(d['profit'],Decimal('800'))
        self.assertEqual(d['equity_start'],Decimal('500'))
        self.assertEqual(d['equity_end'],Decimal('1300'))
        pack=fs.build({2025:self.db},{'years':'2025','basis':'USD'})
        bs=pack['sections'][1]
        assets=next(r for r in bs['rows'] if r[0]=='TOTAL ASSETS')
        self.assertEqual(assets[2],bs['rows'][-1][2])
        self.assertEqual(bs['rows'][-1][0],'TOTAL EQUITY AND LIABILITIES')
        self.assertIn('Single period',str(pack))
        self.assertEqual(pack['sections'][0]['heading'],"INDEPENDENT AUDITOR'S REPORT")
    def test_comparatives_are_separate_and_latest_first(self):
        other=Database(Path(self.folder.name)/'2024.db');other.initialize('secret12345')
        self.voucher(123)
        pack=fs.build({2025:self.db,2024:other},{'years':'2024,2025'})
        self.assertEqual(pack['sections'][1]['headers'],['','Notes','31-12-2025 (USD)','31-12-2024 (USD)'])
        cash=next(r for r in pack['sections'][1]['rows'] if r[0]=='Cash and cash equivalents')
        self.assertEqual(cash[2:],[Decimal('123'),0])
        with self.assertRaisesRegex(ValueError,'not found'): fs.build({2025:self.db},{'years':'2024,2025'})
    def test_saved_text_mapping_currency_and_cash_reconciliation(self):
        self.voucher(100)
        cfg=dict(basis='USD',mapping={'531':'current_assets'},notes={'Entity and activities':'Company-specific text'},audit={'Addressee':'Shareholders'},supplements={'oci':'0','cf_operating':'100','cf_investing':'0','cf_financing':'0','cf_fx':'0'})
        fs.save_config(self.db,cfg,1)
        self.assertEqual(fs.config(self.db),cfg)
        self.assertEqual(fs.year_data(self.db,2025,'USD')['totals']['current_assets'],100)
        self.assertEqual(fs.year_data(self.db,2025,'LBP')['extra'],{})
        cfg['mapping']={};fs.save_config(self.db,cfg,1)
        pack=fs.build({2025:self.db},{'years':[2025]})
        cf=next(s for s in pack['sections'] if s['heading'].startswith('STATEMENT OF CASH FLOWS'))
        self.assertEqual(cf['rows'][-1][2],100)
        self.assertIn('Company-specific text',str(pack))
        self.assertEqual(fs.year_data(self.db,2025,'USD')['profit'],100)
    def test_validation_and_lebanese_classification(self):
        for value in ('', '2024,2025,2026', 'abcd', '2025.5'):
            with self.assertRaises(ValueError): fs.years_from(value)
        for value in ('NaN','Infinity','bad'):
            with self.assertRaises(ValueError): fs.number(value)
        self.assertEqual(fs.default_group('211',1),'intangible')
        self.assertEqual(fs.default_group('221',1),'ppe')
        self.assertEqual(fs.default_group('6511',1),'depreciation')
        self.assertEqual(fs.default_group('6751',1),'finance')
        self.assertEqual(fs.default_group('7751',-1),'finance')
        self.assertEqual(fs.default_group('491',-1),'receivables')
        with self.assertRaises(ValueError): fs.save_config(self.db,{'mapping':{'531':'bad'}},1)
    def test_missing_disclosures_not_assumed_zero_and_loss_sign(self):
        self.voucher(50,debit='6311',credit='531')
        d=fs.year_data(self.db,2025,'USD')
        self.assertEqual(d['profit'],-50)
        pack=fs.build({2025:self.db},{'years':'2025'})
        self.assertIn('PREPARER REVIEW POINTS',str(pack))
        self.assertIn('In our opinion',str(pack))
    def test_exports_include_years_and_disclosures(self):
        from report_export import export_sections_excel,export_sections_pdf
        from openpyxl import load_workbook
        from pypdf import PdfReader
        self.voucher('100.25')
        pack=fs.build({2025:self.db},{'years':'2025'})
        xlsx=Path(self.folder.name)/'pack.xlsx';pdf=Path(self.folder.name)/'pack.pdf'
        export_sections_excel(xlsx,pack['title'],pack['meta'],pack['sections'])
        export_sections_pdf(pdf,pack['title'],pack['meta'],pack['sections'])
        # 2.9.83: the pack is a workbook with a cover and one sheet per statement (it was one long sheet)
        wb=load_workbook(xlsx); self.assertIn('31-12-2025 (USD)',str([list(ws.values) for ws in wb.worksheets]))
        self.assertEqual(wb.sheetnames[:3],['Cover',"Auditor's Report",'Financial Position']);wb.close()
        text=' '.join(p.extract_text() for p in PdfReader(pdf).pages)
        self.assertIn("INDEPENDENT AUDITOR",text);self.assertIn('2025',text);self.assertIn('STATEMENT OF FINANCIAL POSITION',text)


class FinancialStatementsApiTest(unittest.TestCase):
    def test_company_year_routing_and_closed_year_draft_permissions(self):
        import threading
        from http.server import ThreadingHTTPServer
        from server import ApiHandler
        from company_manager import CompanyManager
        from client import ApiClient
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
            path=Path(folder)/'master.db';db=Database(path);db.initialize('secret12345')
            manager=CompanyManager(path)
            handler=type('FinancialTestHandler',(ApiHandler,),dict(db=db,master_db=db,company_manager=manager))
            http=ThreadingHTTPServer(('127.0.0.1',0),handler)
            thread=threading.Thread(target=http.serve_forever,daemon=True);thread.start()
            try:
                api=ApiClient(f'http://127.0.0.1:{http.server_port}');api.login('admin','secret12345')
                company=api.companies()[0];api.select_company_year(company['id'],2024)
                manager.create_year(company['id'],2025,1)
                cfg={'notes':{'Entity and activities':'Year 2025 only'}}
                api.save_financial_config(2025,cfg)
                # 2.9.83 (owner decision): texts entered once apply to every year of the company
                self.assertEqual(api.financial_config(2024),cfg)
                self.assertEqual(api.financial_config(2025),cfg)
                pack=api.business_report('financial_statements',{'years':'2024,2025','basis':'USD'})
                self.assertIn('31-12-2025 (USD)',pack['sections'][1]['headers'])
                with self.assertRaisesRegex(RuntimeError,'Fiscal year not found'):
                    api.business_report('financial_statements',{'years':'2023'})
                registry=manager._read();registry['companies'][0]['years'][0]['status']='closed';manager._write(registry)
                api.save_financial_config(2024,{'notes':{'Entity and activities':'Closed books draft'}})
                db.save_user({'username':'fsviewer','password':'view12345','role':'viewer'},1)
                viewer=ApiClient(api.base_url);viewer.login('fsviewer','view12345');viewer.select_company_year(company['id'],2024)
                with self.assertRaisesRegex(RuntimeError,'read-only'): viewer.save_financial_config(2024,{})
            finally:
                http.shutdown();http.server_close();thread.join(timeout=5)

if __name__=='__main__': unittest.main()
