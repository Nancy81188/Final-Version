"""Official payroll forms as Excel, filled from the company settings and the employee (2.9.39)."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import payroll_excel_forms as forms

COMPANY = {"company_name": "SAMPLE TRADING SARL", "company_mof": "1234567", "company_nssf": "601123",
           "company_address": "Baabda", "company_phone": "01-123456", "company_email": "info@example.com"}
EMPLOYEE = {"employee_number": "100", "full_name": "رامي الخوري", "father_name": "جورج", "mother_name": "ماري",
            "birth_date": "1990-05-04", "birth_place": "بيروت", "nationality": "Lebanese", "national_id": "12345",
            "nssf_number": "98765432", "mof_number": "1234567", "marital_status": "married", "spouse_works": 0,
            "children": 2, "hire_date": "2024-01-15", "job_title": "Accountant", "currency": "LBP", "base_salary": "90000000"}


def values(workbook):
    return [str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value not in (None, "")]


class PayrollExcelFormsTest(unittest.TestCase):
    def test_every_form_is_created_filled_rtl_and_one_printed_page(self):
        with tempfile.TemporaryDirectory() as folder:
            for key, (_title, _builder, needs_employee) in forms.FORMS.items():
                with self.subTest(form=key):
                    path = Path(folder) / f"{key}.xlsx"
                    forms.build_form(key, path, COMPANY, EMPLOYEE if needs_employee else None)
                    workbook = load_workbook(path); text = values(workbook)
                    self.assertIn("SAMPLE TRADING SARL", text)
                    for sheet in workbook.worksheets:
                        self.assertTrue(sheet.sheet_view.rightToLeft)
                        self.assertEqual((sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight), (1, 1))
                    if needs_employee:
                        self.assertTrue("رامي" in text or "رامي الخوري" in text)
                        self.assertIn("جورج", text)

    def test_r3_digits_dates_and_choices(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "r3.xlsx"; forms.build_form("MOF_R3", path, COMPANY, EMPLOYEE)
            sheet = load_workbook(path).active; text = values(load_workbook(path))
            row7 = [sheet.cell(7, c).value for c in range(6, 16)]
            # company MOF number: one digit per box, read left to right (highest column is on the visual left)
            self.assertEqual("".join(str(v) for v in reversed(row7) if v), "1234567")
            self.assertTrue({"04", "05", "1990"} <= set(text))      # birth date day / month / year
            self.assertIn("X", text)                                  # married / spouse not working / monthly pay
            self.assertIn("الخوري", text); self.assertIn("رامي", text)

    def test_employee_forms_need_an_employee_and_nothing_is_invented(self):
        with self.assertRaises(ValueError): forms.build_form("CNSS_41A", None, COMPANY, None)
        data = forms.form_data(COMPANY, {"full_name": "Solo"})
        self.assertEqual((data["first_name"], data["last_name"], data["sex"], data["marital_status"]), ("Solo", "", "", ""))


class EmployeeSexTest(unittest.TestCase):
    def test_sex_is_saved_on_the_employee_and_marks_the_forms(self):
        from database import Database
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as folder:
            db = Database(Path(folder) / "c.db"); db.initialize("secret12345")
            with db.connect() as connection: user = connection.execute("SELECT id FROM users").fetchone()[0]
            saved = db.save_employee({"employee_number": "100000001", "full_name": "Rana Haddad", "currency": "LBP", "base_salary": "1000", "sex": "أنثى"}, user)
            self.assertEqual(saved["sex"], "female")
            again = db.save_employee({**saved, "sex": "male"}, user); self.assertEqual(again["sex"], "male")
            kept = db.save_employee({k: v for k, v in again.items() if k != "sex"}, user); self.assertEqual(kept["sex"], "male")
            data = forms.form_data(COMPANY, saved)
            self.assertEqual(data["sex"], "female")
            path = Path(folder) / "41a.xlsx"; forms.build_form("CNSS_41A", path, COMPANY, saved)
            sheet = load_workbook(path).active
            self.assertIn("X", [str(sheet.cell(14, c).value) for c in range(1, 25)])  # انثى box marked
            if hasattr(db, "release"): db.release()


if __name__ == "__main__":
    unittest.main()
