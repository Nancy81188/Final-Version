"""Inventory screens (version 1.20)."""
from __future__ import annotations
from desktop_common import add_search_bar  # 2.9.78
from desktop_common import search_arrows  # 2.9.78

import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk
from desktop_common import vat_rate, vat_rate_text, vat_currency  # 2.9.72
from desktop_common import account_names, account_label, account_code  # 2.9.79

from desktop_brains import EditableSheet
from multi_select import MultiSelect, chosen_values
from inventory import ANALYSIS_DIMENSIONS
from desktop_common import flow_toolbars

NAVY, GOLD, LIGHT = "#102A43", "#B78B45", "#F4F7FA"  # 2.9.59: the same colours on every screen
RED, MUTED = "#8B1E1E", "#5f6b76"
DOC_TYPES = {"Opening Stock": "opening", "Stock Receipt": "receipt", "Stock Issue": "issue", "Adjustment +": "adjustment_in", "Adjustment -": "adjustment_out", "Transfer": "transfer"}
REPORTS = {"Inventory Summary": "summary", "Stock by Brand & Warehouse": "brands", "Stock Ageing": "ageing", "Inventory Analysis (3D)": "analysis3d",
        "Inventory Health": "health", "Stock Valuation": "valuation", "Stock Card": "stock_card", "Stock Movements": "movements", "Stock Turnover": "turnover",
        "Stock by Supplier": "supplier_stock", "Physical Count Variances": "count_variances", "Sales Margin (COGS)": "margin", "Reorder Report": "reorder",
        "Slow-moving Stock": "slow", "Stock vs Ledger": "ledger_check"}


def _num(value):
    try: return float(str(value or 0).replace(",", ""))
    except ValueError: return None


def _dd(value):
    text = str(value or "")
    return f"{text[8:10]}-{text[5:7]}-{text[:4]}" if len(text) == 10 and text[4] == "-" else text


class InventoryMixin:
    def build_inventory(self):
        nested = ttk.Notebook(self.inventory_tab); nested.pack(fill="both", expand=True, padx=8, pady=8)
        pages = {name: tk.Frame(nested, bg=LIGHT) for name in ("Items", "Categories & Units", "Stock In / Stock Out", "Stock Documents", "Production",
                "Physical Inventory", "Inventory Reports", "Ageing Report", "Warehouses & Settings")}
        for name, page in pages.items(): nested.add(page, text=name)
        self.ageing_tab=pages["Ageing Report"]
        self.build_items_page(pages["Items"]); self.build_categories_page(pages["Categories & Units"]); self.build_stock_in_out_page(pages["Stock In / Stock Out"])
        self.build_physical_page(pages["Physical Inventory"]); self.build_stock_documents_page(pages["Stock Documents"])
        self.build_inventory_reports_page(pages["Inventory Reports"]); self.build_inventory_settings_page(pages["Warehouses & Settings"])
        self.load_inventory()
        self.build_production_page(pages["Production"])  # 2.9.65: recipes, production orders, production report

    def load_inventory(self):
        if not hasattr(self, "items_tree") or not self.items_tree.winfo_exists(): return
        try: self.inventory_rows = self.client.inventory_items(); self.warehouse_rows = self.client.warehouses(); inv = self.client.inventory_settings()
        except Exception as exc: return messagebox.showerror("Inventory", str(exc))
        self.inventory_currency = inv["currency"]
        self.apply_inventory_visibility(inv)
        self.items_tree.delete(*self.items_tree.get_children())
        for i in self.inventory_rows:
            status = "Inactive" if not i["active"] else "Reorder" if i["reorder_level"] and i["quantity"] <= i["reorder_level"] else "OK"
            self.items_tree.insert("", "end", iid=str(i["id"]), values=(i["sku"], i["name"], i.get("category") or "", i.get("subcategory") or "",
                    i.get("brand") or "", i.get("supplier_name") or "", i["unit"],
                f'{i["quantity"]:,.3f}', f'{i["average_cost"]:,.4f}', f'{i["stock_value"]:,.2f}', f'{i["sales_price"]:,.2f}', status), tags=("reorder",) if status == "Reorder" else ())
        self.load_item_lists()
        names = [f'{w["code"]} - {w["name"]}' for w in self.warehouse_rows if w["active"]]
        for box, var in ((getattr(self, "sio_wh_box", None), getattr(self, "sio_vars", {}).get("warehouse")), (getattr(self, "pc_wh_box", None), getattr(self, "pc_vars", {}).get("warehouse"))):
            if box is not None and box.winfo_exists():
                box["values"] = names
                if var is not None and not var.get() and names: var.set(names[0])
        if hasattr(self, "pc_find_box"): self.load_counts()
        for box in (getattr(self, "sd_warehouse_box", None), getattr(self, "sd_to_box", None)):
            if box is not None and box.winfo_exists(): box["values"] = names
        if hasattr(self, "ir_warehouse_box"): self.ir_warehouse_box["values"] = ["All"] + names
        if hasattr(self, "rc_product_box") and self.rc_product_box.winfo_exists():  # 2.9.65: production lists follow the items / warehouses
            self.rc_product_box["values"] = self.item_choices(); self.po_wh_box["values"] = names; self.po_to_box["values"] = names
        items = [f'{i["sku"]} - {i["name"]}' for i in self.inventory_rows]
        for box in (getattr(self, "ir_item_box", None), getattr(self, "ir_item_to_box", None)):
            if box is not None: box["values"] = [""] + items
        categories = sorted({i.get("category") for i in self.inventory_rows if i.get("category")})
        if hasattr(self, "ir_category_box"): self.ir_category_box["values"] = ["All"] + categories
        if hasattr(self, "warehouses_tree"):
            self.warehouses_tree.delete(*self.warehouses_tree.get_children())
            for w in self.warehouse_rows: self.warehouses_tree.insert("", "end", iid=str(w["id"]), values=(w["code"], w["name"], "Yes" if w["active"] else "No"))
            self.inv_currency.set(inv["currency"]); self.inv_method.set("FIFO" if inv["method"] == "fifo" else "Weighted average")
        self.load_stock_documents()

    def item_by_code(self, code):
        code = str(code or "").split(" - ", 1)[0].strip().upper()
        if not code: return None
        return next((i for i in getattr(self, "inventory_rows", []) if i["sku"].upper() == code or (i.get("barcode") or "").upper() == code), None)

    def filter_inventory_find(self, box, variable, map_name):
        from desktop import row_matches_search
        box["values"]=[label for label in getattr(self,map_name, {}) if row_matches_search((label,),variable.get())]

    # ------------------------------------------------------------ items
    def build_items_page(self, page):
        form = tk.LabelFrame(page, text="Item", bg=LIGHT, padx=8, pady=5); form.pack(fill="x", padx=8, pady=6)
        self.item_id = None; self.item_vars = {k: tk.StringVar() for k in ("sku", "name", "unit", "category", "subcategory", "brand", "supplier_name",
                "location", "sales_price", "reorder_level", "barcode", "notes", "default_vat", "cost_account", "stock_account")}
        self.item_vars["unit"].set("unit"); self.item_vars["default_vat"].set(vat_rate_text(self)); self.item_active = tk.BooleanVar(value=True)
        self.item_boxes = {}
        for index, (key, label, width) in enumerate((("sku", "Item Code (auto if blank)", 14), ("name", "Item Name", 24), ("unit", "Unit", 12), ("category", "Category", 16),
                                                     ("subcategory", "Subcategory", 16), ("brand", "Brand", 16), ("supplier_name", "Supplier", 22),
                                                             ("sales_price", "Sales Price", 11), ("default_vat", "Default VAT", 9), ("reorder_level",
                                                             "Reorder Level", 9),
                                                     ("location", "Location (shelf)", 12), ("barcode", "Barcode", 14), ("notes", "Notes", 24),
                                                     ("stock_account", "Stock Account", 26), ("cost_account", "Cost Account (linked)", 24))):
            tk.Label(form, text=label, bg=LIGHT).grid(row=index // 3, column=(index % 3) * 2, sticky="w", padx=4, pady=2)
            if key == "default_vat":
                widget = ttk.Combobox(form, textvariable=self.item_vars[key], values=[vat_rate_text(self), "0%"], state="readonly", width=width); self.item_boxes[key] = widget
            elif key in ("unit", "category", "subcategory", "supplier_name", "brand"):
                widget = ttk.Combobox(form, textvariable=self.item_vars[key], width=width); self.item_boxes[key] = widget
                if key == "category": widget.bind("<<ComboboxSelected>>", lambda _e: self.item_category_chosen())
            elif key == "stock_account":  # 2.9.79: class 3 stock account, linked to its cost account
                widget = ttk.Combobox(form, textvariable=self.item_vars[key], width=width); self.item_boxes[key] = widget
                widget.bind("<<ComboboxSelected>>", lambda _e: self.item_stock_account_chosen(True)); widget.bind("<FocusOut>", lambda _e: self.item_stock_account_chosen(False), add="+")
            elif key == "cost_account": widget = self.account_search_box(form, self.item_vars[key], width)
            else: widget = tk.Entry(form, textvariable=self.item_vars[key], width=width)
            widget.grid(row=index // 3, column=(index % 3) * 2 + 1, sticky="w", padx=4, pady=2)
            if key == "brand": self.item_brand_widgets = (form.grid_slaves(row=index // 3, column=(index % 3) * 2)[0], widget)
        self.item_link_label = tk.Label(form, text="", bg=LIGHT, fg=NAVY, anchor="w", justify="left", wraplength=620); self.item_link_label.grid(row=5, column=0, columnspan=4, sticky="w", padx=4)
        self.item_cost_label = tk.Label(form, text="Cost price (average of purchases): -", bg=LIGHT, fg="#1a5fb4", cursor="hand2", font=("Segoe UI", 9, "bold",
                "underline")); self.item_cost_label.grid(row=5, column=4, columnspan=2, sticky="w", padx=4)
        # The cost is a link: click it to open the item's Stock Card (each purchase / sale and the running average cost).
        self.item_cost_label.bind("<Button-1>", lambda _event: self.open_item_cost_link())
        buttons = tk.Frame(form, bg=LIGHT); buttons.grid(row=6, column=0, columnspan=6, sticky="w", pady=(4, 0))
        self.load_item_stock_accounts(); self.item_vars["stock_account"].set(self.item_stock_choice("37")); self.item_stock_account_chosen(False)
        tk.Checkbutton(buttons, text="Active", variable=self.item_active, bg=LIGHT).pack(side="left", padx=(0, 8))
        self.action_button(buttons, "New", self.new_item).pack(side="left", padx=3)
        tk.Button(buttons, text="Save", command=self.save_item, bg=GOLD, fg=NAVY, border=0, padx=18, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        self.action_button(buttons, "Stock Card", lambda: self.open_stock_card()).pack(side="left", padx=3)
        tk.Label(buttons, text="Quantity and cost come from the stock documents. Double-click an item to edit it.", bg=LIGHT, fg=MUTED).pack(side="left", padx=10)
        self.items_tree = self.table(page, [("sku", "Item Code", 95), ("name", "Item", 200), ("category", "Category", 95), ("subcategory", "Subcategory", 95),
                ("brand", "Brand", 90), ("supplier", "Supplier", 120),
            ("unit", "Unit", 50), ("qty", "On Hand", 80), ("cost", "Cost Price (avg)", 100), ("value", "Stock Value", 100), ("price", "Sales Price", 85), ("status", "Status", 65)])
        self.items_tree.tag_configure("reorder", foreground=RED); self.items_tree.bind("<Double-1>", lambda _e: self.edit_item())

    # 2.9.79: stock account (class 3) linked to the cost account and the year-end variation accounts
    def load_item_stock_accounts(self):
        names = account_names(self)
        codes = sorted(code for code in names if code.isdigit() and code.startswith("3") and not code.startswith("39") and len(code) >= 2)
        self._item_stock_choices = [f"{code} - {names[code]}" for code in codes] or ["37 - Stock of Goods for Sale"]
        box = getattr(self, "item_boxes", {}).get("stock_account")
        if box is not None: box["values"] = self._item_stock_choices

    def item_stock_choice(self, code):
        code = account_code(code) or "37"
        return next((label for label in getattr(self, "_item_stock_choices", []) if account_code(label) == code), account_label(self, code))

    def item_stock_account_chosen(self, fill_cost):
        """Show the link of the chosen stock account; when it is chosen from the list, put its linked cost account in Cost Account."""
        import inventory as inventory_rules
        code = account_code(self.item_vars["stock_account"].get())
        try: cost, opening, closing, note = inventory_rules.stock_link(code or "37")
        except ValueError as exc:
            if hasattr(self, "item_link_label"): self.item_link_label.config(text=str(exc), fg=RED)
            return
        if fill_cost:
            current = account_code(self.item_vars["cost_account"].get())
            if cost and (not current or current in {link[0] for link in inventory_rules.STOCK_LINKS.values()}): self.item_vars["cost_account"].set(account_label(self, cost))
            elif not cost and current in {link[0] for link in inventory_rules.STOCK_LINKS.values() if link[0]}: self.item_vars["cost_account"].set("")
        year_end = f"Dr {opening} / Cr {code}, then Dr {code} / Cr {closing}" if opening != closing else f"variation on {account_label(self, opening)}"
        if hasattr(self, "item_link_label"):
            self.item_link_label.config(text=f"Link: {note}. Purchases: Cost Account ({'empty = purchase screen account' if code.startswith('37') or not cost else 'empty = ' + cost}). "
                                         f"Year end: {year_end}.", fg=NAVY)

    def new_item(self):
        self.item_id = None; [v.set("") for v in self.item_vars.values()]; self.item_vars["unit"].set("unit"); self.item_vars["default_vat"].set(vat_rate_text(self)); self.item_active.set(True)
        self.item_vars["stock_account"].set(self.item_stock_choice("37")); self.item_stock_account_chosen(False)
        if hasattr(self, "item_cost_label"): self.item_cost_label.config(text="Cost price (average of purchases): -")

    def edit_item(self):
        selected = self.items_tree.selection()
        if not selected: return
        item = next(i for i in self.inventory_rows if str(i["id"]) == selected[0]); self.item_id = item["id"]
        for key in self.item_vars: self.item_vars[key].set("" if item.get(key) in (None, 0.0) and key in ("barcode", "notes", "category",
                "brand") else str(item.get(key) if item.get(key) is not None else ""))
        self.item_vars["default_vat"].set("0%" if str(item.get("default_vat") or "11").strip() in ("0", "0.0", "0%") else vat_rate_text(self))
        self.item_vars["sales_price"].set(f'{item["sales_price"]:g}'); self.item_vars["reorder_level"].set(f'{item["reorder_level"]:g}'); self.item_active.set(bool(item["active"]))
        self.item_vars["stock_account"].set(self.item_stock_choice(item.get("stock_account") or "37"))
        self.item_vars["cost_account"].set(account_label(self, item.get("cost_account") or "")); self.item_stock_account_chosen(False)
        self.item_cost_label.config(text=f'Cost price (average of purchases): {item["average_cost"]:,.4f} {getattr(self, "inventory_currency", "")}   On hand: {item["quantity"]:,.3f}   ▸ how is it calculated?')

    def save_item(self):
        payload = {k: v.get().strip() for k, v in self.item_vars.items()}; payload.update(id=self.item_id, active=self.item_active.get())
        payload["default_vat"] = "0" if payload.get("default_vat", "11%").replace("%", "").strip() in ("0", "0.0") else f"{vat_rate(self):g}"
        try: saved = self.client.save_inventory_item(payload)
        except Exception as exc: return messagebox.showerror("Items", str(exc))
        self.new_item(); self.load_inventory(); messagebox.showinfo("Items", f'Item {saved["sku"]} - {saved["name"]} saved')

    def open_item_cost_link(self):
        if not self.items_tree.selection(): return messagebox.showinfo("Cost price", "Select an item in the list, then click its cost to see how it is calculated (Stock Card).")
        self.open_stock_card()

    def open_stock_card(self):
        selected = self.items_tree.selection()
        if not selected: return messagebox.showwarning("Stock Card", "Select an item first")
        item = next(i for i in self.inventory_rows if str(i["id"]) == selected[0])
        self.ir_report.set("Stock Card"); self.ir_item.set(f'{item["sku"]} - {item["name"]}'); self.ir_item_to.set("")
        self.inventory_report_selected()
        self.inventory_notebook_select("Inventory Reports"); self.run_inventory_report()

    def inventory_notebook_select(self, name):
        notebook = [w for w in self.inventory_tab.winfo_children() if isinstance(w, ttk.Notebook)][0]
        notebook.select([t for t in notebook.tabs() if notebook.tab(t, "text") == name][0])

    # ------------------------------------------------------------ stock documents
    def build_stock_documents_page(self, page):
        self.sd_id = None; self.sd_vars = {k: tk.StringVar() for k in ("type", "number", "date", "warehouse", "to_warehouse", "party", "reference", "notes", "find", "project", "branch")}
        v = self.sd_vars; v["type"].set("Stock Receipt"); v["date"].set(self.fiscal_today())
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=(6, 2))
        tk.Label(bar, text="Type", bg=LIGHT, font=("Segoe UI", 9, "bold")).pack(side="left")
        type_box = ttk.Combobox(bar, textvariable=v["type"], values=list(DOC_TYPES), state="readonly", width=14); type_box.pack(side="left", padx=(4, 8))
        type_box.bind("<<ComboboxSelected>>", lambda _e: self.stock_type_changed())
        tk.Label(bar, text="Number", bg=LIGHT).pack(side="left")
        tk.Entry(bar, textvariable=v["number"], width=16, state="readonly", readonlybackground="white", font=("Segoe UI", 10, "bold")).pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Date", bg=LIGHT).pack(side="left"); self.date_entry(bar, v["date"], 11).pack(side="left", padx=(4, 8))
        self.sd_warehouse_label = tk.Label(bar, text="Warehouse", bg=LIGHT); self.sd_warehouse_label.pack(side="left")
        self.sd_warehouse_box = ttk.Combobox(bar, textvariable=v["warehouse"], state="readonly", width=17); self.sd_warehouse_box.pack(side="left", padx=(4, 8))
        self.sd_to_label = tk.Label(bar, text="To", bg=LIGHT); self.sd_to_box = ttk.Combobox(bar, textvariable=v["to_warehouse"], state="readonly", width=17)
        tk.Label(bar, text="Find", bg=LIGHT).pack(side="right")
        self.sd_find_box = ttk.Combobox(bar, textvariable=v["find"], width=26); self.sd_find_box.pack(side="right", padx=4)
        self.sd_find_box.bind("<<ComboboxSelected>>", lambda _e: self.open_stock_document())
        self.sd_find_box.bind("<KeyRelease>", lambda _e: self.filter_inventory_find(self.sd_find_box, self.sd_vars["find"], "sd_doc_map"))
        self.sd_find_box.bind("<Return>", lambda _e: self.open_stock_document())
        bar2 = tk.Frame(page, bg=LIGHT); bar2.pack(fill="x", padx=8, pady=2)
        tk.Label(bar2, text="Customer / Supplier", bg=LIGHT).pack(side="left")
        self.sd_party_box = ttk.Combobox(bar2, textvariable=v["party"], width=26); self.sd_party_box.pack(side="left", padx=(4, 8))
        tk.Label(bar2, text="Reference (invoice / PO)", bg=LIGHT).pack(side="left"); tk.Entry(bar2, textvariable=v["reference"], width=16).pack(side="left", padx=(4, 8))
        tk.Label(bar2, text="Notes", bg=LIGHT).pack(side="left"); tk.Entry(bar2, textvariable=v["notes"], width=28).pack(side="left", padx=4)
        self.sd_project_group = tk.Frame(bar2, bg=LIGHT); self.sd_project_group.pack(side="left", padx=(8, 0))
        tk.Label(self.sd_project_group, text="Project", bg=LIGHT).pack(side="left")
        self.sd_project_box = ttk.Combobox(self.sd_project_group, textvariable=v["project"], state="readonly", width=18); self.sd_project_box.pack(side="left", padx=4)
        self.sd_branch_group = tk.Frame(bar2, bg=LIGHT); self.sd_branch_group.pack(side="left", padx=(8, 0))
        tk.Label(self.sd_branch_group, text="Branch", bg=LIGHT).pack(side="left")
        self.sd_branch_box = ttk.Combobox(self.sd_branch_group, textvariable=v["branch"], state="readonly", width=16); self.sd_branch_box.pack(side="left", padx=4)
        bottom = tk.Frame(page, bg=LIGHT); bottom.pack(side="bottom", fill="x", padx=8, pady=6)
        self.action_button(bottom, "New", self.new_stock_document).pack(side="left", padx=(0, 3))
        self.action_button(bottom, "New Warehouse Transfer", self.new_warehouse_transfer).pack(side="left", padx=3)
        self.action_button(bottom, "Add Line", self.add_stock_line).pack(side="left", padx=3)
        tk.Button(bottom, text="Delete Line", command=self.delete_stock_line, bg=RED, fg="white", border=0, padx=12, pady=7).pack(side="left", padx=3)
        tk.Button(bottom, text="Save", command=self.save_stock_document, bg=GOLD, fg=NAVY, border=0, padx=18, pady=7, font=("Segoe UI", 9, "bold")).pack(side="left", padx=(12, 3))
        tk.Button(bottom, text="Delete Document", command=self.delete_stock_document, bg=RED, fg="white", border=0, padx=12, pady=7).pack(side="left", padx=3)
        self.sd_total = tk.Label(bottom, text="", bg=LIGHT, fg=NAVY, font=("Segoe UI", 10, "bold")); self.sd_total.pack(side="right", padx=8)
        self.sd_info = tk.Label(page, text="Item Code: type the code or press F2 in the cell. Issues, adjustments - and transfers leave at the current cost; receipts and opening stock need a unit cost.",
                                bg=LIGHT, fg=MUTED, anchor="w"); self.sd_info.pack(side="bottom", fill="x", padx=10)
        columns = [("line", "#", 40, "center"), ("sku", "Item Code", 110, "w"), ("name", "Item", 260, "w"), ("unit", "Unit", 60, "center"), ("quantity", "Quantity", 100, "e"),
                   ("unit_cost", "Unit Cost", 110, "e"), ("value", "Value", 120, "e"), ("on_hand", "On Hand", 100, "e")]
        self.stock_sheet = EditableSheet(self, page, columns, ["sku", "quantity", "unit_cost"], self.stock_cell_changed, height=9)
        self.stock_sheet.tree.bind("<F2>", lambda _e: self.stock_item_lookup())
        self.new_stock_document()

    def stock_item_lookup(self):
        iid, row = self.stock_sheet.selected()
        if not row: return
        window = tk.Toplevel(self); window.title("Items - F2"); window.geometry("640x420"); window.configure(bg=LIGHT); window.transient(self); window.grab_set()
        search = tk.StringVar(); entry = tk.Entry(window, textvariable=search, width=40); entry.pack(padx=10, pady=8); entry.focus_set()
        tree = ttk.Treeview(window, columns=("sku", "name", "qty"), show="headings"); tree.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        for key, label, width in (("sku", "Code", 110), ("name", "Item", 320), ("qty", "On Hand", 100)): tree.heading(key, text=label); tree.column(key, width=width)
        def fill(*_a):
            tree.delete(*tree.get_children()); text = search.get().casefold()
            for i in self.inventory_rows:
                if i["active"] and (not text or text in f'{i["sku"]} {i["name"]}'.casefold()): tree.insert("", "end", values=(i["sku"], i["name"], f'{i["quantity"]:,.3f}'))
        def choose(_e=None):
            if tree.selection(): self.stock_cell_changed(iid, "sku", tree.item(tree.selection()[0], "values")[0]); self.stock_sheet.refresh(iid); window.destroy()
        search.trace_add("write", fill); tree.bind("<Double-1>", choose); tree.bind("<Return>", choose); fill(); search_arrows(entry, tree, choose, search)  # 2.9.78

    def stock_type_changed(self):
        transfer = self.sd_vars["type"].get() == "Transfer"
        self.sd_warehouse_label.config(text="From" if transfer else "Warehouse")
        if transfer: self.sd_to_label.pack(side="left"); self.sd_to_box.pack(side="left", padx=(4, 8))
        else: self.sd_to_label.pack_forget(); self.sd_to_box.pack_forget()
        self.sd_info.config(text=("Choose different From and To warehouses. Stock moves at its existing cost; the On Hand column is company-wide, but saving checks stock in From."
                                  if transfer else "Item Code: type the code or press F2 in the cell. Issues and adjustments - use automatic cost; receipts and opening stock need a unit cost."))
        if not self.sd_id:
            try: self.sd_vars["number"].set(self.client.next_stock_number(DOC_TYPES[self.sd_vars["type"].get()], self.sd_vars["date"].get()))
            except Exception: self.sd_vars["number"].set("")
        for iid, row in self.stock_sheet.rows.items(): self.refresh_stock_row(row); self.stock_sheet.refresh(iid)
        self.update_stock_total()

    def cost_is_automatic(self):
        return DOC_TYPES[self.sd_vars["type"].get()] in ("issue", "adjustment_out", "transfer")

    def refresh_stock_row(self, row):
        item = self.item_by_code(row.get("sku"))
        if item:
            row["name"] = item["name"]; row["unit"] = item["unit"]; row["on_hand"] = f'{item["quantity"]:,.3f}'
            if self.cost_is_automatic() or not row.get("unit_cost"): row["unit_cost"] = round(item["average_cost"], 4) if self.cost_is_automatic() or item["average_cost"] else row.get("unit_cost", "")
        qty = _num(row.get("quantity")) or 0; cost = _num(row.get("unit_cost")) or 0
        row["value"] = f"{qty * cost:,.2f}" if qty else ""
        row["_display"] = {"quantity": f"{qty:,.3f}" if qty else "", "unit_cost": f"{cost:,.4f}" if row.get("unit_cost") not in ("", None) else ""}

    def stock_cell_changed(self, iid, key, text):
        row = self.stock_sheet.rows[iid]
        if key == "sku":
            item = self.item_by_code(text)
            if text and not item: messagebox.showwarning("Stock Documents", f"Item {text} was not found. Press F2 on the line to search."); return False
            row["sku"] = item["sku"] if item else ""
        else:
            value = _num(text)
            if value is None or value < 0: messagebox.showwarning("Stock Documents", "Enter a positive number"); return False
            if key == "unit_cost" and self.cost_is_automatic(): messagebox.showinfo("Stock Documents", "Issues, adjustments - and transfers use the current average cost automatically"); return False
            row[key] = value
        self.refresh_stock_row(row); self.update_stock_total()
        rows = self.stock_sheet.tree.get_children()
        if rows and iid == rows[-1] and row.get("sku") and row.get("quantity"): self.add_stock_line(edit=False)

    def add_stock_line(self, edit=True):
        row = {"sku": "", "quantity": "", "unit_cost": ""}; self.refresh_stock_row(row); iid = self.stock_sheet.insert(row)
        self.stock_sheet.tree.selection_set(iid); self.stock_sheet.tree.focus(iid)
        if edit: self.after(30, lambda: self.stock_sheet.edit(iid, "sku"))

    def delete_stock_line(self):
        if not self.stock_sheet.delete_selected(): return messagebox.showwarning("Stock Documents", "Select a line first")
        self.update_stock_total()

    def update_stock_total(self):
        rows = [r for r in self.stock_sheet.ordered() if r.get("sku") and _num(r.get("quantity"))]
        total = sum((_num(r["quantity"]) or 0) * (_num(r.get("unit_cost")) or 0) for r in rows)
        label = "Estimated value (average-cost preview)" if self.sd_vars["type"].get() == "Transfer" else "Total value"
        self.sd_total.config(text=f"{len(rows)} line(s)   {label}: {total:,.2f} {getattr(self, 'inventory_currency', '')}")

    def new_stock_document(self):
        self.sd_id = None; v = self.sd_vars
        for key in ("party", "reference", "notes", "find"): v[key].set("")
        v["date"].set(self.fiscal_today())
        if not v["warehouse"].get() and getattr(self, "warehouse_rows", None): v["warehouse"].set(f'{self.warehouse_rows[0]["code"]} - {self.warehouse_rows[0]["name"]}')
        self.stock_sheet.clear(); self.add_stock_line(edit=False); self.add_stock_line(edit=False); self.stock_type_changed()

    def new_warehouse_transfer(self):
        active = [f'{w["code"]} - {w["name"]}' for w in getattr(self, "warehouse_rows", []) if w["active"]]
        if len(active) < 2:
            return messagebox.showwarning("Warehouse Transfer", "Create two active warehouses under Warehouses & Settings first.")
        self.inventory_notebook_select("Stock Documents")
        self.sd_vars["type"].set("Transfer")
        self.new_stock_document()
        source = self.sd_vars["warehouse"].get()
        if source not in active:
            source = active[0]; self.sd_vars["warehouse"].set(source)
        self.sd_vars["to_warehouse"].set(next(name for name in active if name != source))

    def load_stock_documents(self):
        if not hasattr(self, "sd_find_box"): return
        try: docs = self.client.stock_documents(); parties = self.client.parties()
        except Exception: return
        labels = {v: k for k, v in DOC_TYPES.items()}
        self.sd_doc_map = {f'{d["number"]} | {_dd(d["doc_date"])} | {labels[d["doc_type"]]} | {d.get("party_name") or d.get("reference") or ""}': d["id"]
                           for d in docs if d["doc_type"] in labels}  # production orders are opened from Inventory > Production
        self.sd_find_box["values"] = list(self.sd_doc_map)
        self.sd_party_map = {p["name"]: p for p in parties}; self.sd_party_box["values"] = list(self.sd_party_map)

    def open_stock_document(self):
        document_id = getattr(self, "sd_doc_map", {}).get(self.sd_vars["find"].get())
        if not document_id:
            matches=list(self.sd_find_box["values"])
            if len(matches)==1: self.sd_vars["find"].set(matches[0]); document_id=self.sd_doc_map.get(matches[0])
        if not document_id: return
        try: doc = self.client.stock_document(document_id)
        except Exception as exc: return messagebox.showerror("Stock Documents", str(exc))
        self.sd_id = doc["id"]; v = self.sd_vars; labels = {val: k for k, val in DOC_TYPES.items()}
        v["type"].set(labels[doc["doc_type"]]); v["number"].set(doc["number"]); v["date"].set(_dd(doc["doc_date"]))
        v["project"].set(f'{doc["project_code"]} - {doc["project_name"]}' if doc.get("project_code") else ""); v["branch"].set(doc.get("branch_name") or "")
        v["warehouse"].set(f'{doc["warehouse_code"]} - {doc["warehouse_name"]}'); v["party"].set(doc.get("party_name") or ""); v["reference"].set(doc.get("reference") or ""); v["notes"].set(doc.get("notes") or "")
        if doc.get("to_warehouse_code"): v["to_warehouse"].set(next((f'{w["code"]} - {w["name"]}' for w in self.warehouse_rows if w["code"] == doc["to_warehouse_code"]), ""))
        self.stock_sheet.clear()
        for line in doc["lines"]:
            row = {"sku": line["sku"], "quantity": line["quantity"], "unit_cost": line["unit_cost"]}; self.refresh_stock_row(row); self.stock_sheet.insert(row)
        self.stock_type_changed(); v["number"].set(doc["number"])
        self.sd_info.config(text=f"Document {doc['number']} opened" + (" (created by an invoice - change the invoice instead)" if doc.get("invoice_id") else ""))

    def save_stock_document(self):
        v = self.sd_vars; lines = [r for r in self.stock_sheet.ordered() if r.get("sku") and _num(r.get("quantity"))]
        if not lines: return messagebox.showwarning("Stock Documents", "Add at least one item with a quantity")
        doc_type = DOC_TYPES[v["type"].get()]
        if doc_type == "transfer":
            source = v["warehouse"].get().split(" - ", 1)[0]
            target = v["to_warehouse"].get().split(" - ", 1)[0]
            if not source or not target or source == target:
                return messagebox.showwarning("Warehouse Transfer", "Choose different From and To warehouses.")
        if doc_type in ("opening", "receipt") and any(not _num(r.get("unit_cost")) for r in lines): return messagebox.showwarning("Stock Documents", "Enter the unit cost of every line")
        party = getattr(self, "sd_party_map", {}).get(v["party"].get())
        header = {"doc_type": doc_type, "doc_date": v["date"].get().strip(), "warehouse_id": v["warehouse"].get().split(" - ", 1)[0], "to_warehouse_id": v["to_warehouse"].get().split(" - ", 1)[0],
                  "party_id": party["id"] if party else None, "reference": v["reference"].get().strip(), "notes": v["notes"].get().strip(),
                  "project_id": v["project"].get().split(" - ", 1)[0].strip(), "branch_id": v["branch"].get().strip()}
        payload = [{"sku": r["sku"], "quantity": r["quantity"], "unit_cost": 0 if self.cost_is_automatic() else r.get("unit_cost")} for r in lines]
        try: saved = self.client.save_stock_document(header, payload, self.sd_id)
        except Exception as exc: return messagebox.showerror("Stock Documents", str(exc))
        messagebox.showinfo("Stock Documents", f'{saved["number"]} saved'); self.load_inventory(); self.new_stock_document()

    def delete_stock_document(self):
        if not self.sd_id: return messagebox.showwarning("Stock Documents", "Open a saved document first (Find)")
        if not messagebox.askyesno("Stock Documents", f"Delete {self.sd_vars['number'].get()}?"): return
        try: self.client.delete_stock_document(self.sd_id)
        except Exception as exc: return messagebox.showerror("Stock Documents", str(exc))
        self.load_inventory(); self.new_stock_document()

    # ------------------------------------------------------------ reports
    def build_inventory_reports_page(self, page):
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=6); year = getattr(self, "current_fiscal_year", datetime.now().year)
        self.ir_report = tk.StringVar(value="Stock Valuation"); self.ir_from = tk.StringVar(value=f"01-01-{year}"); self.ir_to = tk.StringVar(value=f"31-12-{year}")
        self.ir_warehouse = tk.StringVar(value="All"); self.ir_item = tk.StringVar(); self.ir_item_to = tk.StringVar(); self.ir_method = tk.StringVar(value="Company setting"); self.ir_days = tk.StringVar(value="90")
        self.ir_category = tk.StringVar(value="All"); self.ir_zero = tk.BooleanVar(value=False)
        self.ir_report_box = ttk.Combobox(bar, textvariable=self.ir_report, values=list(REPORTS), state="readonly", width=22)
        self.ir_report_box.pack(side="left", padx=(0, 8))
        self.ir_report_box.bind("<<ComboboxSelected>>", self.inventory_report_selected)
        tk.Label(bar, text="From Date", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.ir_from, 11).pack(side="left", padx=(4, 6))
        tk.Label(bar, text="To Date / As of", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.ir_to, 11).pack(side="left", padx=(4, 6))
        self.ir_warehouse_label = tk.Label(bar, text="Warehouse", bg=LIGHT); self.ir_warehouse_label.pack(side="left")
        self.ir_warehouse_box = MultiSelect(bar, self.ir_warehouse, width=16, title="Warehouse", bg=LIGHT); self.ir_warehouse_box.pack(side="left", padx=(4, 6))
        tk.Label(bar, text="Costing", bg=LIGHT).pack(side="left")
        ttk.Combobox(bar, textvariable=self.ir_method, values=["Company setting", "Weighted average", "FIFO"], state="readonly", width=15).pack(side="left", padx=4)
        item_row = tk.Frame(page, bg=LIGHT); self.ir_item_row = item_row
        tk.Label(item_row, text="Item From", bg=LIGHT, font=("Segoe UI", 9, "bold")).pack(side="left")
        self.ir_item_box = ttk.Combobox(item_row, textvariable=self.ir_item, width=29); self.ir_item_box.pack(side="left", padx=(4, 12))
        tk.Label(item_row, text="Item To", bg=LIGHT, font=("Segoe UI", 9, "bold")).pack(side="left")
        self.ir_item_to_box = ttk.Combobox(item_row, textvariable=self.ir_item_to, width=29); self.ir_item_to_box.pack(side="left", padx=(4, 8))
        for widget, variable in ((self.ir_item_box, self.ir_item), (self.ir_item_to_box, self.ir_item_to)):
            widget.bind("<KeyRelease>", lambda event, box=widget, var=variable: self.filter_report_items(event, box, var))
            widget.bind("<Return>", lambda event, box=widget, var=variable: self.select_report_item(event, box, var))
        bar2 = tk.Frame(page, bg=LIGHT); bar2.pack(fill="x", padx=8); self.ir_controls_row = bar2
        tk.Label(bar2, text="Category", bg=LIGHT).pack(side="left")
        self.ir_category_box = MultiSelect(bar2, self.ir_category, width=14, title="Category", bg=LIGHT); self.ir_category_box.pack(side="left", padx=(4, 8))
        self.ir_subcategory = tk.StringVar(value="All"); self.ir_unit = tk.StringVar(value="All"); self.ir_supplier = tk.StringVar(value="All")
        bar3 = tk.Frame(page, bg=LIGHT); bar3.pack(fill="x", padx=8, pady=(4, 0), after=bar2)
        tk.Label(bar3, text="Subcategory", bg=LIGHT).pack(side="left"); self.ir_subcategory_box = MultiSelect(bar3, self.ir_subcategory, width=14,
                title="Subcategory", bg=LIGHT); self.ir_subcategory_box.pack(side="left", padx=(4, 8))
        tk.Label(bar3, text="Unit", bg=LIGHT).pack(side="left"); self.ir_unit_box = MultiSelect(bar3, self.ir_unit, width=8, title="Unit", bg=LIGHT); self.ir_unit_box.pack(side="left", padx=(4, 8))
        tk.Label(bar3, text="Supplier", bg=LIGHT).pack(side="left"); self.ir_supplier_box = MultiSelect(bar3, self.ir_supplier, width=22, title="Supplier",
                bg=LIGHT); self.ir_supplier_box.pack(side="left", padx=(4, 8))
        self.ir_brand = tk.StringVar(value="All"); self.ir_project = tk.StringVar(value="All"); self.ir_branch = tk.StringVar(value="All"); self.ir_filter_groups = {}
        for key, label, variable, width in (("brand", "Brand", self.ir_brand, 14), ("project", "Project", self.ir_project, 18), ("branch", "Branch", self.ir_branch, 14)):
            group = tk.Frame(bar3, bg=LIGHT); group.pack(side="left"); self.ir_filter_groups[key] = group
            tk.Label(group, text=label, bg=LIGHT).pack(side="left")
            box = MultiSelect(group, variable, width=width, title=label, bg=LIGHT); box.pack(side="left", padx=(4, 8)); setattr(self, f"ir_{key}_box", box)
        self.inventory_report_selected()
        tk.Label(bar2, text="Slow-moving days", bg=LIGHT).pack(side="left"); tk.Entry(bar2, textvariable=self.ir_days, width=5).pack(side="left", padx=4)
        tk.Checkbutton(bar2, text="Include zero stock", variable=self.ir_zero, bg=LIGHT).pack(side="left", padx=6)
        tk.Button(bar2, text="Show", command=self.run_inventory_report, bg=GOLD, fg=NAVY, border=0, padx=18, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=6)
        for text, fmt in (("Print", "print"), ("Excel", "xlsx"), ("PDF", "pdf")): self.action_button(bar2, text, lambda f=fmt: self.export_inventory_report(f)).pack(side="left", padx=2)
        self.ir_info = tk.Label(page, text="", bg=LIGHT, fg=NAVY, anchor="w"); self.ir_info.pack(fill="x", padx=10)
        # Inventory Analysis (3D) options: shown only when that report is selected.
        analysis_bar = tk.Frame(page, bg=LIGHT); self.ir_analysis_bar = analysis_bar
        self.ir_3d_rows = tk.StringVar(value="Item"); self.ir_3d_columns = tk.StringVar(value="Warehouse"); self.ir_3d_layers = tk.StringVar(value="(none)")
        self.ir_3d_measure = tk.StringVar(value="Quantity")
        dimensions = [d.title() for d in ANALYSIS_DIMENSIONS]
        for label, variable, values in (("Rows", self.ir_3d_rows, dimensions), ("Columns", self.ir_3d_columns, dimensions),
                                        ("Layers (3rd)", self.ir_3d_layers, ["(none)"] + dimensions), ("Measure", self.ir_3d_measure, ["Quantity", "Value"])):
            tk.Label(analysis_bar, text=label, bg=LIGHT, font=("Segoe UI", 9, "bold")).pack(side="left")
            ttk.Combobox(analysis_bar, textvariable=variable, values=values, state="readonly", width=12).pack(side="left", padx=(4, 10))
        self.action_button(analysis_bar, "3D Chart", self.show_inventory_3d_chart).pack(side="left", padx=4)
        tk.Label(analysis_bar, text="Choose any 3. Stock at To date; with Month, Project or Branch: net movement From / To.", bg=LIGHT, fg=MUTED).pack(side="left", padx=6)
        flow_toolbars(bar2, analysis_bar)  # 2.9.59 (bar / bar3 hide filters by setting, they keep pack)
        self.inventory_report_selected()
        self.ir_viewer = self.report_viewer(page)

    def filter_report_items(self, event=None, box=None, variable=None):
        from desktop import row_matches_search
        if event is not None and event.keysym in ("Up", "Down", "Return", "Escape", "Tab"): return
        box = box or self.ir_item_box; variable = variable or self.ir_item
        choices=[f'{item["sku"]} - {item["name"]}' for item in getattr(self,"inventory_rows",[])]
        box["values"]=[choice for choice in choices if row_matches_search((choice,),variable.get())]

    def select_report_item(self, _event=None, box=None, variable=None):
        box = box or self.ir_item_box; variable = variable or self.ir_item
        matches=list(box["values"])
        if len(matches)==1: variable.set(matches[0])

    def inventory_report_selected(self, _event=None):
        self.ir_item_row.pack(fill="x", padx=8, pady=(2, 0), before=self.ir_controls_row)
        if hasattr(self, "ir_analysis_bar"):
            if self.ir_report.get() == "Inventory Analysis (3D)": self.ir_analysis_bar.pack(fill="x", padx=8, pady=(2, 0), before=self.ir_info)
            else: self.ir_analysis_bar.pack_forget()

    def inventory_report_options(self):
        options = {"date_from": self.ir_from.get().strip(), "date_to": self.ir_to.get().strip(), "days": self.ir_days.get().strip() or "90", "include_zero": self.ir_zero.get()}
        if self.ir_report.get() == "Inventory Analysis (3D)":
            layers = self.ir_3d_layers.get().lower() if hasattr(self, "ir_3d_layers") else "none"
            options.update(rows=self.ir_3d_rows.get().lower(), columns=self.ir_3d_columns.get().lower(), measure=self.ir_3d_measure.get().lower(),
                           layers="none" if layers in ("", "(none)") else layers)
        chosen = [value.split(" - ", 1)[0] for value in chosen_values(self.ir_warehouse.get())]
        ids = [w["id"] for w in getattr(self, "warehouse_rows", []) if w["code"] in chosen]
        if ids: options["warehouse_id"] = ids[0] if len(ids) == 1 else ids
        if self.ir_item.get().strip():
            item = self.item_by_code(self.ir_item.get())
            if not item: raise ValueError("Choose a valid Item From")
            options["item_id"] = item["id"]
        if self.ir_item_to.get().strip():
            to_item = self.item_by_code(self.ir_item_to.get())
            if not to_item: raise ValueError("Choose a valid Item To")
            if not options.get("item_id"): raise ValueError("Choose Item From before Item To")
            options["item_to_id"] = to_item["id"]
        # 2.9.50: every filter below can hold several values (any combination); one value is sent as before
        def pick(values): return values[0] if len(values) == 1 else values
        for key, variable in (("category", self.ir_category), ("subcategory", self.ir_subcategory), ("unit", self.ir_unit), ("brand", getattr(self, "ir_brand", None))):
            values = chosen_values(variable.get()) if variable is not None else []
            if values: options[key] = pick(values)
        if getattr(self, "ir_project", None):
            projects = [p for p in getattr(self, "inventory_projects", []) if any(v.startswith(f'{p["code"]} - ') or v == p["code"] for v in chosen_values(self.ir_project.get()))]
            if projects: options["project_id"] = pick([p["id"] for p in projects]); options["project_name"] = ", ".join(p["name"] for p in projects)
        if getattr(self, "ir_branch", None):
            wanted = {v.casefold() for v in chosen_values(self.ir_branch.get())}
            branches = [b for b in getattr(self, "inventory_branches", []) if b["name"].casefold() in wanted]
            if branches: options["branch_id"] = pick([b["id"] for b in branches]); options["branch_name"] = ", ".join(b["name"] for b in branches)
        suppliers = [getattr(self, "ir_supplier_map", {}).get(name) for name in chosen_values(self.ir_supplier.get())]
        suppliers = [s for s in suppliers if s]
        if suppliers: options["supplier_id"] = pick([s["id"] for s in suppliers]); options["supplier_name"] = pick([s["name"] for s in suppliers])
        method = {"Weighted average": "average", "FIFO": "fifo"}.get(self.ir_method.get())
        if method: options["method"] = method
        return options

    def run_inventory_report(self):
        try:
            start=datetime.strptime(self.ir_from.get().strip(), "%d-%m-%Y")
            end=datetime.strptime(self.ir_to.get().strip(), "%d-%m-%Y")
            if start>end: raise ValueError("From Date must be on or before To Date")
        except ValueError as exc: return messagebox.showwarning("Inventory Reports",str(exc) if "From Date" in str(exc) else "Enter From Date and To Date as DD-MM-YYYY")
        if REPORTS[self.ir_report.get()] == "stock_card" and not self.item_by_code(self.ir_item.get()):
            return messagebox.showwarning("Stock Card", "Choose an item from the list before showing its stock card")
        try: result = self.client.inventory_report(REPORTS[self.ir_report.get()], self.inventory_report_options())
        except Exception as exc: return messagebox.showerror("Inventory Reports", str(exc))
        self.inventory_report_result = result; self.show_sections(self.ir_viewer, result["sections"]); self.ir_info.config(text=f'{result["title"]}  |  ' + "   ".join(result["meta"]))

    def show_inventory_3d_chart(self):
        """The 3D chart of the Inventory Analysis (3D) report, with angle and depth controls."""
        if self.ir_report.get() != "Inventory Analysis (3D)": self.ir_report.set("Inventory Analysis (3D)"); self.inventory_report_selected()
        self.run_inventory_report()
        result = getattr(self, "inventory_report_result", None)
        charts = [(s["chart"]["title"], s["chart"]) for s in (result or {}).get("sections", []) if s.get("chart") and s["chart"].get("series")]
        if not charts: return messagebox.showinfo("3D Chart", "No data to draw for these choices")
        from desktop_projection import show_3d_charts
        return show_3d_charts(self, charts, "Inventory Analysis (3D)")

    def export_inventory_report(self, format_name):
        if not getattr(self, "inventory_report_result", None): self.run_inventory_report()
        result = getattr(self, "inventory_report_result", None)
        if result: self.output_sections(result["title"], result["meta"], result["sections"], result["title"].replace(" ", "_"), format_name)

    # ------------------------------------------------------------ warehouses & settings
    def build_inventory_settings_page(self, page):
        box = tk.LabelFrame(page, text="Inventory settings", bg=LIGHT, padx=8, pady=6); box.pack(fill="x", padx=8, pady=6)
        self.inv_currency = tk.StringVar(value="USD"); self.inv_method = tk.StringVar(value="Weighted average")
        tk.Label(box, text="Stock valued in", bg=LIGHT).pack(side="left"); ttk.Combobox(box, textvariable=self.inv_currency, values=self.currency_codes,
                state="readonly", width=6).pack(side="left", padx=(4, 10))
        tk.Label(box, text="Costing method", bg=LIGHT).pack(side="left"); ttk.Combobox(box, textvariable=self.inv_method, values=["Weighted average", "FIFO"],
                state="readonly", width=16).pack(side="left", padx=(4, 10))
        self.inv_show = {key: tk.BooleanVar(value=True) for key in ("brand", "warehouse", "project", "branch")}
        show_box = tk.LabelFrame(page, text="Show in the inventory screens | يظهر في الشاشات", bg=LIGHT, padx=8, pady=4); show_box.pack(fill="x", padx=8, pady=(0, 6))
        for key, label in (("brand", "Brand"), ("warehouse", "Warehouse"), ("project", "Project"), ("branch", "Branch")):
            tk.Checkbutton(show_box, text=label, variable=self.inv_show[key], bg=LIGHT).pack(side="left", padx=8)
        tk.Label(show_box, text="Tick only what you use; the others disappear from items, stock documents and report filters. Save Settings to keep the choice.",
                bg=LIGHT, fg=MUTED).pack(side="left", padx=8)
        self.inv_ask = {key: tk.BooleanVar(value=False) for key in ("category", "subcategory", "brand")}  # 2.9.101
        ask_box = tk.LabelFrame(page, text="New item on a purchase: ask for | صنف جديد في المشتريات: اطلب", bg=LIGHT, padx=8, pady=4); ask_box.pack(fill="x", padx=8, pady=(0, 6))
        for key, label in (("category", "Category"), ("subcategory", "Subcategory"), ("brand", "Brand")):
            tk.Checkbutton(ask_box, text=label, variable=self.inv_ask[key], bg=LIGHT).pack(side="left", padx=8)
        tk.Label(ask_box, bg=LIGHT, fg=MUTED,
                 text="When a purchase creates new items, a table opens to choose their category / brand. Save Settings to keep the choice.").pack(side="left", padx=8)
        self.action_button(box, "Save Settings", self.save_inventory_settings).pack(side="left", padx=4)
        year = getattr(self, "current_fiscal_year", datetime.now().year)
        tk.Button(box, text=f"Post Stock Variation {year}", command=self.post_stock_variation, bg=GOLD, fg=NAVY, border=0, padx=14, pady=6, font=("Segoe UI", 9,
                "bold")).pack(side="left", padx=(20, 4))
        # 2.9.82: the stock variation of a month, so the P&L shows each month's cost of sales and gross margin
        self.monthly_variation_end = tk.StringVar(value=datetime.now().strftime("%m-%Y") if str(datetime.now().year) == str(year) else f"12-{year}")
        tk.Label(box, text="Month (MM-YYYY)", bg=LIGHT).pack(side="left", padx=(12, 2)); tk.Entry(box, textvariable=self.monthly_variation_end, width=8).pack(side="left")
        self.action_button(box, "Post Monthly Variation", self.post_monthly_stock_variation).pack(side="left", padx=4)
        tk.Label(page, text="Lebanese periodic method: purchases stay in 601 / 611. The Stock Variation voucher (type 06) cancels the stock in each item's stock account and books the counted "
                 "closing stock: goods 37 against 6051 / 6052, raw materials 31 against 6151 / 6152, work in progress 33 against 7211, products 35 against 7255 (the Stock Account of the item). "
                 "It is made automatically when you close the year, and the closing stock becomes the Opening Stock of the next year.",
                 bg=LIGHT, fg=MUTED, wraplength=1080, justify="left").pack(fill="x", padx=12, pady=(0, 6))
        wh = tk.LabelFrame(page, text="Warehouses", bg=LIGHT, padx=8, pady=6); wh.pack(fill="both", expand=True, padx=8, pady=6)
        self.wh_id = None; self.wh_code = tk.StringVar(); self.wh_name = tk.StringVar(); self.wh_active = tk.BooleanVar(value=True)
        form = tk.Frame(wh, bg=LIGHT); form.pack(fill="x")
        tk.Label(form, text="Code (blank = automatic)", bg=LIGHT).pack(side="left"); tk.Entry(form, textvariable=self.wh_code, width=8).pack(side="left", padx=(4, 10))
        tk.Label(form, text="Name", bg=LIGHT).pack(side="left"); tk.Entry(form, textvariable=self.wh_name, width=28).pack(side="left", padx=(4, 10))
        tk.Checkbutton(form, text="Active", variable=self.wh_active, bg=LIGHT).pack(side="left")
        self.action_button(form, "New", lambda: (setattr(self, "wh_id", None), self.wh_code.set(""), self.wh_name.set(""), self.wh_active.set(True))).pack(side="left", padx=3)
        tk.Button(form, text="Save Warehouse", command=self.save_warehouse, bg=GOLD, fg=NAVY, border=0, padx=14, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        self.warehouses_tree = ttk.Treeview(wh, columns=("code", "name", "active"), show="headings", height=6)
        for key, label, width in (("code", "Code", 90), ("name", "Warehouse", 300), ("active", "Active", 80)): self.warehouses_tree.heading(key,
                text=label); self.warehouses_tree.column(key, width=width)
        self.warehouses_tree.pack(fill="both", expand=True, pady=6); self.warehouses_tree.bind("<Double-1>", lambda _e: self.edit_warehouse()); add_search_bar(self.warehouses_tree)  # 2.9.78

    def edit_warehouse(self):
        selected = self.warehouses_tree.selection()
        if not selected: return
        w = next(w for w in self.warehouse_rows if str(w["id"]) == selected[0]); self.wh_id = w["id"]; self.wh_code.set(w["code"]); self.wh_name.set(w["name"]); self.wh_active.set(bool(w["active"]))

    def save_warehouse(self):
        try: self.client.save_warehouse({"id": self.wh_id, "code": self.wh_code.get(), "name": self.wh_name.get(), "active": self.wh_active.get()})
        except Exception as exc: return messagebox.showerror("Warehouses", str(exc))
        self.wh_id = None; self.wh_code.set(""); self.wh_name.set(""); self.load_inventory()

    def save_inventory_settings(self):
        payload = {"currency": self.inv_currency.get(), "method": "fifo" if self.inv_method.get() == "FIFO" else "average"}
        payload.update({f"show_{key}": var.get() for key, var in getattr(self, "inv_show", {}).items()})
        payload.update({f"ask_{key}": var.get() for key, var in getattr(self, "inv_ask", {}).items()})
        try: self.client.save_inventory_settings(payload)
        except Exception as exc: return messagebox.showerror("Inventory", str(exc))
        self.load_inventory(); messagebox.showinfo("Inventory", "Inventory settings saved")

    def post_stock_variation(self):
        year = getattr(self, "current_fiscal_year", datetime.now().year)
        if not messagebox.askyesno("Stock Variation", f"Post (or replace) the Stock Variation voucher of {year} with the stock value at 31-12-{year}?"): return
        try: result = self.client.post_stock_variation(year)
        except Exception as exc: return messagebox.showerror("Stock Variation", str(exc))
        lines = [f"Voucher {result.get('voucher') or '-'}", f"Stock in the ledger before: {result['opening']:,.2f}", f"Closing stock: {result['closing']:,.2f}"]
        for group in result.get("groups") or []:  # 2.9.79: one line per stock account
            lines.append(f"{account_label(self, group['stock_account'])}: before {group['opening']:,.2f} (Dr {group['opening_account']}), closing {group['closing']:,.2f} (Cr {group['closing_account']})")
        messagebox.showinfo("Stock Variation", "\n".join(lines))
        self.load_journal(); self.load_trial()

    def post_monthly_stock_variation(self):
        import calendar
        text = self.monthly_variation_end.get().strip()
        try: month, year = (int(part) for part in text.split("-")); last = f"{calendar.monthrange(year, month)[1]:02d}-{month:02d}-{year}"
        except (ValueError, calendar.IllegalMonthError): return messagebox.showwarning("Stock Variation", "Enter the month as MM-YYYY, for example 03-2026")
        if not messagebox.askyesno("Stock Variation", f"Post (or replace) the stock variation of {month:02d}-{year} at {last}?\nLater months already posted are posted again, in order."): return
        try: result = self.client.post_monthly_stock_variation(last)
        except Exception as exc: return messagebox.showerror("Stock Variation", str(exc))
        lines = [f'{p["date"]}: voucher {p["voucher"]}, stock before {p["opening"]:,.2f}, at month end {p["closing"]:,.2f}' for p in result.get("posted") or []]
        messagebox.showinfo("Stock Variation", "\n".join(lines) or "Nothing to post: the ledger already shows the stock of that month.")
        self.load_journal(); self.load_trial()

    # ------------------------------------------------------------ lists for the item form and filters
    def apply_inventory_visibility(self, inv):
        """Brand / warehouse / project / branch appear only when ticked in Warehouses & Settings."""
        show = {key: inv.get(f"show_{key}", True) for key in ("brand", "warehouse", "project", "branch")}
        for key, var in getattr(self, "inv_show", {}).items(): var.set(show[key])
        for key, var in getattr(self, "inv_ask", {}).items(): var.set(bool(inv.get(f"ask_{key}")))
        def toggle(widget, visible, manager="pack", **options):
            if widget is None or not widget.winfo_exists(): return
            if visible:
                if not widget.winfo_manager(): (widget.grid(**options) if manager == "grid" else widget.pack(**options))
            else: (widget.grid_remove() if widget.winfo_manager() == "grid" else widget.pack_forget())
        for widget in getattr(self, "item_brand_widgets", ()):
            if not show["brand"]: widget.grid_remove()
            else: widget.grid()
        if not show["brand"]: self.items_tree.configure(displaycolumns=[c for c in self.items_tree["columns"] if c != "brand"])
        else: self.items_tree.configure(displaycolumns="#all")
        for key, var in (("project", self.sd_vars.get("project")), ("branch", self.sd_vars.get("branch"))):
            toggle(getattr(self, f"sd_{key}_group", None), show[key], side="left", padx=(8, 0))
            if not show[key] and var is not None: var.set("")
        for key in ("brand", "project", "branch"):
            group = getattr(self, "ir_filter_groups", {}).get(key)
            toggle(group, show[key], side="left")
            if not show[key] and hasattr(self, f"ir_{key}"): getattr(self, f"ir_{key}").set("All")
        for widget in (getattr(self, "ir_warehouse_label", None), getattr(self, "ir_warehouse_box", None)):
            if widget is not None and not show["warehouse"]: widget.pack_forget(); self.ir_warehouse.set("All")
        self.inventory_visibility = show

    def load_item_lists(self):
        try: data = self.client.item_categories(); parties = self.client.parties()
        except Exception: return
        try:
            brands = self.client.inventory_brands(); self.inventory_projects = self.client.projects(); self.inventory_branches = self.client.branches()
        except Exception: brands = []; self.inventory_projects = []; self.inventory_branches = []
        if hasattr(self, "item_boxes") and "brand" in self.item_boxes: self.item_boxes["brand"]["values"] = brands
        projects = [f'{p["code"]} - {p["name"]}' for p in self.inventory_projects if p.get("active", 1)]
        branches = [b["name"] for b in self.inventory_branches if b.get("active", 1)]
        for name, values in (("ir_brand_box", ["All"] + brands), ("ir_project_box", ["All"] + projects), ("ir_branch_box", ["All"] + branches),
                             ("sd_project_box", [""] + projects), ("sd_branch_box", [""] + branches)):
            box = getattr(self, name, None)
            if box is not None and box.winfo_exists(): box["values"] = values
        self.category_data = data; suppliers = [p for p in parties if p["kind"] in ("supplier", "both")]
        self.ir_supplier_map = {p["name"]: p for p in suppliers}
        categories = [c["name"] for c in data["categories"]]; subcategories = sorted({s for c in data["categories"] for s in c["subcategories"]})
        if hasattr(self, "item_boxes"):
            self.item_boxes["category"]["values"] = categories; self.item_boxes["unit"]["values"] = data["units"]; self.item_boxes["supplier_name"]["values"] = [p["name"] for p in suppliers]
            self.item_category_chosen()
        for name, values in (("ir_category_box", ["All"] + categories), ("ir_subcategory_box", ["All"] + subcategories), ("ir_unit_box", ["All"] + data["units"]),
                             ("ir_supplier_box", ["All"] + [p["name"] for p in suppliers])):
            box = getattr(self, name, None)
            if box is not None and box.winfo_exists(): box["values"] = values
        if hasattr(self, "cat_tree") and self.cat_tree.winfo_exists():
            self.cat_tree.delete(*self.cat_tree.get_children())
            for c in data["categories"]:
                parent = self.cat_tree.insert("", "end", text=c["name"], values=("Category",), open=True)
                for sub in c["subcategories"]: self.cat_tree.insert(parent, "end", text=sub, values=("Subcategory",))
            self.unit_list.delete(0, "end"); [self.unit_list.insert("end", u) for u in data["units"]]
            self.cat_parent_box["values"] = categories

    def item_category_chosen(self):
        data = getattr(self, "category_data", None)
        if not data or not hasattr(self, "item_boxes"): return
        chosen = next((c for c in data["categories"] if c["name"] == self.item_vars["category"].get()), None)
        self.item_boxes["subcategory"]["values"] = chosen["subcategories"] if chosen else []

    # ------------------------------------------------------------ categories, subcategories, units
    def build_categories_page(self, page):
        left = tk.LabelFrame(page, text="Categories and subcategories", bg=LIGHT, padx=8, pady=6); left.pack(side="left", fill="both", expand=True, padx=8, pady=8)
        self.cat_name = tk.StringVar(); self.cat_parent = tk.StringVar(); self.unit_name = tk.StringVar()
        form = tk.Frame(left, bg=LIGHT); form.pack(fill="x")
        tk.Label(form, text="New category", bg=LIGHT).grid(row=0, column=0, sticky="w"); tk.Entry(form, textvariable=self.cat_name, width=22).grid(row=0, column=1, padx=4)
        self.action_button(form, "Add Category", lambda: self.save_category("category")).grid(row=0, column=2, padx=4)
        self.sub_name = tk.StringVar()
        tk.Label(form, text="New subcategory", bg=LIGHT).grid(row=1, column=0, sticky="w", pady=4); tk.Entry(form, textvariable=self.sub_name, width=22).grid(row=1, column=1, padx=4)
        tk.Label(form, text="in", bg=LIGHT).grid(row=1, column=2, sticky="w"); self.cat_parent_box = ttk.Combobox(form, textvariable=self.cat_parent,
                state="readonly", width=18); self.cat_parent_box.grid(row=1, column=3, padx=4)
        self.action_button(form, "Add Subcategory", lambda: self.save_category("subcategory")).grid(row=1, column=4, padx=4)
        self.cat_tree = ttk.Treeview(left, columns=("level",), height=14); self.cat_tree.heading("#0", text="Name"); self.cat_tree.heading("level", text="Level")
        self.cat_tree.column("#0", width=260); self.cat_tree.pack(fill="both", expand=True, pady=6)
        right = tk.LabelFrame(page, text="Units", bg=LIGHT, padx=8, pady=6); right.pack(side="left", fill="y", padx=8, pady=8)
        row = tk.Frame(right, bg=LIGHT); row.pack(fill="x")
        tk.Entry(row, textvariable=self.unit_name, width=14).pack(side="left"); self.action_button(row, "Add Unit", lambda: self.save_category("unit")).pack(side="left", padx=4)
        self.unit_list = tk.Listbox(right, height=16, width=22); self.unit_list.pack(fill="y", expand=True, pady=6)

    def save_category(self, kind):
        name = {"category": self.cat_name, "subcategory": self.sub_name, "unit": self.unit_name}[kind].get().strip()
        try: self.client.save_item_category({"kind": kind, "name": name, "parent": self.cat_parent.get()})
        except Exception as exc: return messagebox.showerror("Categories & Units", str(exc))
        {"category": self.cat_name, "subcategory": self.sub_name, "unit": self.unit_name}[kind].set(""); self.load_item_lists()

    # ------------------------------------------------------------ stock in / stock out
    def build_stock_in_out_page(self, page):
        self.sio_vars = {k: tk.StringVar() for k in ("direction", "date", "warehouse", "reason")}
        self.sio_vars["direction"].set("Stock In (add)"); self.sio_vars["date"].set(self.fiscal_today())
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=6)
        ttk.Combobox(bar, textvariable=self.sio_vars["direction"], values=["Stock In (add)", "Stock Out (remove)"], state="readonly", width=18).pack(side="left", padx=(0, 8))
        tk.Label(bar, text="Date", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.sio_vars["date"], 11).pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Warehouse", bg=LIGHT).pack(side="left"); self.sio_wh_box = ttk.Combobox(bar, textvariable=self.sio_vars["warehouse"],
                state="readonly", width=18); self.sio_wh_box.pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Reason", bg=LIGHT).pack(side="left"); tk.Entry(bar, textvariable=self.sio_vars["reason"], width=28).pack(side="left", padx=4)
        self.action_button(bar, "Transfer Between Warehouses", self.new_warehouse_transfer).pack(side="right", padx=4)
        bottom = tk.Frame(page, bg=LIGHT); bottom.pack(side="bottom", fill="x", padx=8, pady=6)
        self.action_button(bottom, "Add Line", lambda: self.sio_sheet.insert(self.sio_row())).pack(side="left", padx=(0, 3))
        tk.Button(bottom, text="Delete Line", command=lambda: self.sio_sheet.delete_selected(), bg=RED, fg="white", border=0, padx=10, pady=6).pack(side="left", padx=3)
        tk.Button(bottom, text="Save", command=self.save_stock_in_out, bg=GOLD, fg=NAVY, border=0, padx=18, pady=6, font=("Segoe UI", 9, "bold")).pack(side="left", padx=(12, 3))
        tk.Label(bottom, text="Stock In: cost = average cost unless you type another. Stock Out: leaves at the average cost. Quantities and costs update at once.",
                bg=LIGHT, fg=MUTED).pack(side="left", padx=8)
        self.sio_sheet = EditableSheet(self, page, [("line", "#", 40, "center"), ("sku", "Item Code", 110, "w"), ("name", "Item", 260, "w"), ("unit", "Unit", 60, "center"),
            ("on_hand", "On Hand", 100, "e"), ("quantity", "Quantity", 100, "e"), ("unit_cost", "Unit Cost (average)", 130, "e"), ("value", "Value", 120, "e")],
            ["sku", "quantity", "unit_cost"], self.sio_changed, height=10)
        self.sio_sheet.tree.bind("<F2>", lambda _e: self.item_picker(lambda sku: self.sio_pick(sku)))
        for _ in range(3): self.sio_sheet.insert(self.sio_row())

    def sio_row(self):
        return {"sku": "", "name": "", "unit": "", "on_hand": "", "quantity": "", "unit_cost": "", "value": ""}

    def sio_pick(self, sku):
        iid, _row = self.sio_sheet.selected()
        if iid: self.sio_changed(iid, "sku", sku); self.sio_sheet.refresh(iid)

    def sio_changed(self, iid, key, text):
        row = self.sio_sheet.rows[iid]
        if key == "sku":
            item = self.item_by_code(text)
            if text and not item: messagebox.showwarning("Stock In / Out", f"Item {text} was not found (F2 to search)"); return False
            if item: row.update(sku=item["sku"], name=item["name"], unit=item["unit"], on_hand=f'{item["quantity"]:,.3f}', unit_cost=round(item["average_cost"], 4))
        else:
            value = _num(text)
            if value is None or value < 0: messagebox.showwarning("Stock In / Out", "Enter a positive number"); return False
            if key == "unit_cost" and self.sio_vars["direction"].get().startswith("Stock Out"): messagebox.showinfo("Stock In / Out", "Stock out leaves at the average cost"); return False
            row[key] = value
        row["value"] = f'{(_num(row.get("quantity")) or 0) * (_num(row.get("unit_cost")) or 0):,.2f}'

    def save_stock_in_out(self):
        lines = [r for r in self.sio_sheet.ordered() if r.get("sku") and _num(r.get("quantity"))]
        if not lines: return messagebox.showwarning("Stock In / Out", "Add at least one item with a quantity")
        stock_in = self.sio_vars["direction"].get().startswith("Stock In")
        header = {"doc_type": "adjustment_in" if stock_in else "adjustment_out", "doc_date": self.sio_vars["date"].get(),
                "warehouse_id": (self.sio_vars["warehouse"].get() or "MAIN").split(" - ", 1)[0],
                  "notes": self.sio_vars["reason"].get().strip() or ("Stock in" if stock_in else "Stock out")}
        try: saved = self.client.save_stock_document(header, [{"sku": r["sku"], "quantity": r["quantity"], "unit_cost": r.get("unit_cost") if stock_in else 0} for r in lines])
        except Exception as exc: return messagebox.showerror("Stock In / Out", str(exc))
        messagebox.showinfo("Stock In / Out", f'{saved["number"]} saved'); self.sio_sheet.clear(); [self.sio_sheet.insert(self.sio_row()) for _ in range(3)]; self.load_inventory()

    def item_picker(self, callback):
        window = tk.Toplevel(self); window.title("Items - F2"); window.geometry("660x420"); window.configure(bg=LIGHT); window.transient(self); window.grab_set()
        search = tk.StringVar(); entry = tk.Entry(window, textvariable=search, width=40); entry.pack(padx=10, pady=8); entry.focus_set()
        tree = ttk.Treeview(window, columns=("sku", "name", "category", "qty"), show="headings"); tree.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        for key, label, width in (("sku", "Code", 100), ("name", "Item", 280), ("category", "Category", 120), ("qty", "On Hand", 90)): tree.heading(key, text=label); tree.column(key, width=width)
        def fill(*_a):
            tree.delete(*tree.get_children()); text = search.get().casefold()
            for i in getattr(self, "inventory_rows", []):
                if i["active"] and (not text or text in f'{i["sku"]} {i["name"]} {i.get("category") or ""} {i.get("barcode") or ""}'.casefold()): tree.insert("",
                        "end", values=(i["sku"], i["name"], i.get("category") or "", f'{i["quantity"]:,.3f}'))
        def choose(_e=None):
            if tree.selection(): value = tree.item(tree.selection()[0], "values")[0]; window.destroy(); callback(value)
        search.trace_add("write", fill); tree.bind("<Double-1>", choose); tree.bind("<Return>", choose); fill(); search_arrows(entry, tree, choose, search)  # 2.9.78

    # ------------------------------------------------------------ physical inventory
    def build_physical_page(self, page):
        self.pc_id = None; self.pc_vars = {k: tk.StringVar() for k in ("date", "warehouse", "find")}; self.pc_vars["date"].set(self.fiscal_today())
        bar = tk.Frame(page, bg=LIGHT); bar.pack(fill="x", padx=8, pady=6)
        tk.Label(bar, text="Count date", bg=LIGHT).pack(side="left"); self.date_entry(bar, self.pc_vars["date"], 11).pack(side="left", padx=(4, 8))
        tk.Label(bar, text="Warehouse", bg=LIGHT).pack(side="left"); self.pc_wh_box = ttk.Combobox(bar, textvariable=self.pc_vars["warehouse"],
                state="readonly", width=18); self.pc_wh_box.pack(side="left", padx=(4, 8))
        tk.Button(bar, text="Load Stock on Hand", command=self.load_count_sheet, bg=GOLD, fg=NAVY, border=0, padx=12, pady=5, font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        self.action_button(bar, "Add Item", self.add_physical_item).pack(side="left", padx=3)
        tk.Label(bar, text="Open count", bg=LIGHT).pack(side="left", padx=(12, 2)); self.pc_find_box = ttk.Combobox(bar, textvariable=self.pc_vars["find"],
                state="readonly", width=34); self.pc_find_box.pack(side="left")
        self.pc_find_box.bind("<<ComboboxSelected>>", lambda _e: self.open_count())
        self.pc_find_box.bind("<KeyRelease>", lambda _e: self.filter_inventory_find(self.pc_find_box, self.pc_vars["find"], "pc_map"))
        self.pc_find_box.bind("<Return>", lambda _e: self.open_count())
        bottom = tk.Frame(page, bg=LIGHT); bottom.pack(side="bottom", fill="x", padx=8, pady=6)
        self.action_button(bottom, "Save Count", lambda: self.save_count(False)).pack(side="left", padx=(0, 3))
        tk.Button(bottom, text="Apply Count to Stock + Journal", command=lambda: self.save_count(True), bg=GOLD, fg=NAVY, border=0, padx=14, pady=6,
                font=("Segoe UI", 9, "bold")).pack(side="left", padx=3)
        self.action_button(bottom, "Print / PDF", lambda: self.count_document("preview")).pack(side="left", padx=3)
        self.action_button(bottom, "Excel Count Sheet", self.export_count_excel).pack(side="left", padx=3)
        self.action_button(bottom, "Upload Counts (Excel)", self.upload_count_excel).pack(side="left", padx=3)
        self.pc_info = tk.Label(bottom, text="", bg=LIGHT, fg=NAVY, font=("Segoe UI", 9, "bold")); self.pc_info.pack(side="right", padx=6)
        self.pc_sheet = EditableSheet(self, page, [("line", "#", 40, "center"), ("sku", "Item Code", 100, "w"), ("name", "Item", 240, "w"), ("location", "Location", 80, "center"),
            ("unit", "Unit", 55, "center"), ("system_qty", "Stock on Hand", 110, "e"), ("counted", "Physical Count", 110, "e"), ("difference", "Difference", 100, "e"),
            ("value", "Difference Value", 120, "e")], ["counted"], self.count_changed, height=12)

    def load_count_sheet(self):
        warehouse = self.warehouse_id_of(self.pc_vars["warehouse"].get())
        try: rows = self.client.count_sheet(warehouse, self.pc_vars["date"].get())
        except Exception as exc: return messagebox.showerror("Physical Inventory", str(exc))
        self.pc_id = None; self.pc_sheet.clear()
        for row in rows: row["counted"] = ""; self.count_display(row); self.pc_sheet.insert(row)
        self.update_count_info()

    def add_physical_item(self):
        def add(sku):
            item = self.item_by_code(sku)
            if not item: return
            if any(row["item_id"] == item["id"] for row in self.pc_sheet.ordered()): return messagebox.showinfo("Physical Inventory", "Item already on the count sheet")
            try: rows = self.client.count_sheet(self.warehouse_id_of(self.pc_vars["warehouse"].get()), self.pc_vars["date"].get())
            except Exception as exc: return messagebox.showerror("Physical Inventory", str(exc))
            row = next((r for r in rows if r["item_id"] == item["id"]), None)
            if row: row["counted"] = ""; self.count_display(row); self.pc_sheet.insert(row); self.update_count_info()
        self.item_picker(add)

    def warehouse_id_of(self, label):
        code = (label or "MAIN").split(" - ", 1)[0]
        return next((w["id"] for w in getattr(self, "warehouse_rows", []) if w["code"] == code), 1)

    def count_display(self, row):
        counted = _num(row.get("counted")) if row.get("counted") not in ("", None) else None
        difference = (counted - row["system_qty"]) if counted is not None else None
        row["difference"] = difference; row["_display"] = {"system_qty": f'{row["system_qty"]:,.3f}', "counted": f"{counted:,.3f}" if counted is not None else "",
            "difference": f"{difference:+,.3f}" if difference else ("0" if difference == 0 else ""), "value": f'{difference * row["unit_cost"]:+,.2f}' if difference else ""}

    def count_changed(self, iid, key, text):
        row = self.pc_sheet.rows[iid]
        if text.strip() == "": row["counted"] = ""
        else:
            value = _num(text)
            if value is None or value < 0: messagebox.showwarning("Physical Inventory", "Enter the counted quantity (0 or more)"); return False
            row["counted"] = value
        self.count_display(row); self.update_count_info()

    def update_count_info(self):
        rows = self.pc_sheet.ordered(); counted = [r for r in rows if r.get("counted") not in ("", None)]
        value = sum((r["difference"] or 0) * r["unit_cost"] for r in counted)
        self.pc_info.config(text=f"{len(counted)} of {len(rows)} item(s) counted   Difference value: {value:+,.2f} {getattr(self, 'inventory_currency', '')}")

    def save_count(self, post):
        lines = [{"item_id": r["item_id"], "sku": r["sku"], "counted": r["counted"]} for r in self.pc_sheet.ordered() if r.get("counted") not in ("", None)]
        if not lines: return messagebox.showwarning("Physical Inventory", "Enter at least one physical count")
        if post and not messagebox.askyesno("Physical Inventory", "Apply the physical count?\n\n- The stock quantities become the counted quantities (adjustment + / - documents).\n"
                                            "- The difference is posted in the journal (JV): surplus Dr stock 37 / Cr stock variation, shortage the other way,\n  valued at the item cost on the count date."): return
        try: saved = self.client.save_physical_count({"count_date": self.pc_vars["date"].get(), "warehouse_id": self.warehouse_id_of(self.pc_vars["warehouse"].get())}, lines, self.pc_id, post)
        except Exception as exc: return messagebox.showerror("Physical Inventory", str(exc))
        self.pc_id = saved["id"]
        messagebox.showinfo("Physical Inventory",
                f'{saved["number"]} saved' + (f' and applied: {saved.get("adjustment_numbers") or "no difference"}\n(the JV is in the General Journal)' if post else " (draft)"))
        self.load_inventory(); self.load_counts()
        if post:
            for refresh in ("load_journal", "load_trial"):
                try: getattr(self, refresh)()
                except Exception: pass  # the screen refreshes when it is opened

    def load_counts(self):
        try: counts = self.client.physical_counts()
        except Exception: counts = []
        self.pc_map = {f'{c["number"]} | {_dd(c["count_date"])} | {c["warehouse_code"]} | {c["status"]}': c["id"] for c in counts}; self.pc_find_box["values"] = list(self.pc_map)

    def open_count(self):
        count_id = getattr(self, "pc_map", {}).get(self.pc_vars["find"].get())
        if not count_id:
            matches=list(self.pc_find_box["values"])
            if len(matches)==1: self.pc_vars["find"].set(matches[0]); count_id=self.pc_map.get(matches[0])
        if not count_id: return
        try: count = self.client.physical_count(count_id); rows = self.client.count_sheet(count["warehouse_id"], count["count_date"])
        except Exception as exc: return messagebox.showerror("Physical Inventory", str(exc))
        counted = {l["item_id"]: l["counted"] for l in count["lines"]}; self.pc_id = None if count["status"] == "posted" else count["id"]; self.pc_sheet.clear()
        self.pc_vars["date"].set(_dd(count["count_date"]))
        for row in rows: row["counted"] = float(counted[row["item_id"]]) if row["item_id"] in counted else ""; self.count_display(row); self.pc_sheet.insert(row)
        self.update_count_info()
        if count["status"] == "posted": self.pc_info.config(text=f'{count["number"]} is posted ({count.get("adjustment_numbers") or "no difference"}): shown for printing')

    def count_rows_for_export(self):
        return [[r["sku"], r["name"], r.get("location") or "", r["unit"], r["system_qty"], r["counted"] if r.get("counted") not in ("", None) else "",
                r["difference"] if r.get("difference") is not None else "",
                 round((r["difference"] or 0) * r["unit_cost"], 2) if r.get("difference") else ""] for r in self.pc_sheet.ordered()]

    def count_document(self, mode):
        rows = self.count_rows_for_export()
        if not rows: return messagebox.showwarning("Physical Inventory", "Load the stock on hand first")
        self.output_sections("Physical Inventory Count", [f"Date: {self.pc_vars['date'].get()}   Warehouse: {self.pc_vars['warehouse'].get()}",
                "Counted by: ______________   Checked by: ______________"],
            [{"heading": "Count sheet", "headers": ["Item Code", "Item", "Location", "Unit", "Stock on Hand", "Physical Count", "Difference",
                    "Difference Value"], "rows": rows, "total_rows": []}], "Physical_Count", mode)

    def export_count_excel(self):
        from report_export import export_excel
        rows = self.count_rows_for_export()
        if not rows: return messagebox.showwarning("Physical Inventory", "Load the stock on hand first")
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="Physical_Count.xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path: return
        export_excel(path, "Physical Count", ["Item Code", "Item", "Location", "Unit", "Stock on Hand", "Physical Count", "Difference", "Difference Value"], rows)
        messagebox.showinfo("Physical Inventory", f"Saved: {path}\nFill the 'Physical Count' column and use 'Upload Counts (Excel)'.")

    def upload_count_excel(self):
        from openpyxl import load_workbook
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if not path: return
        if not self.pc_sheet.ordered(): self.load_count_sheet()
        workbook = load_workbook(path, read_only=True, data_only=True); found = 0
        try:
            rows = list(workbook.worksheets[0].iter_rows(values_only=True))
            header_index = next(i for i, r in enumerate(rows[:10]) if r and any(str(v or "").strip().lower() in ("item code", "code", "sku") for v in r))
            header = [str(v or "").strip().lower() for v in rows[header_index]]
            code_col = next(i for i, h in enumerate(header) if h in ("item code", "code", "sku")); count_col = next(i for i, h in enumerate(header) if "count" in h or h in ("qty", "quantity"))
            counts = {str(r[code_col]).strip().upper(): r[count_col] for r in rows[header_index + 1:] if r and r[code_col] not in (None, "") and r[count_col] not in (None, "")}
        except StopIteration: return messagebox.showerror("Physical Inventory", "The Excel file needs an 'Item Code' column and a 'Physical Count' column")
        finally: workbook.close()
        for iid, row in self.pc_sheet.rows.items():
            if row["sku"].upper() in counts:
                try: row["counted"] = float(str(counts[row["sku"].upper()]).replace(",", "")); found += 1
                except ValueError: continue
                self.count_display(row); self.pc_sheet.refresh(iid)
        self.update_count_info(); messagebox.showinfo("Physical Inventory", f"{found} count(s) loaded from the Excel file. Check them, then Save or Post.")
